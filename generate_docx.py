import base64
import json
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_solution_document(output_docx_path: str):
    doc = docx.Document()

    # Set page margins to 1 inch (72 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    COLOR_PRIMARY = RGBColor(31, 78, 121)     # Navy #1F4E79
    COLOR_DARK = RGBColor(33, 37, 41)          # #212529
    COLOR_MUTED = RGBColor(85, 85, 85)         # #555555

    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_table_borders(table, color="D0D5DD"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_instruction_block(text):
        for line in text.strip().split("\n"):
            line_str = line.strip()
            if not line_str:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line_str)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = COLOR_MUTED

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY

    def add_result_table(headers, rows, right_align_indices=None, col_widths=None):
        if right_align_indices is None:
            right_align_indices = []
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        # Header row
        hdr_cells = table.rows[0].cells
        for col_idx, h_text in enumerate(headers):
            cell = hdr_cells[col_idx]
            cell.text = str(h_text)
            set_cell_background(cell, "1F4E79")
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in right_align_indices else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)

        # Data rows
        for row_idx, row_data in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            bg_color = "F8F9FA" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, val in enumerate(row_data):
                cell = row_cells[col_idx]
                cell.text = str(val)
                if bg_color != "FFFFFF":
                    set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_idx in right_align_indices else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)
                    r.font.color.rgb = COLOR_DARK

        # Set column widths if provided
        if col_widths:
            for row in table.rows:
                for idx, w in enumerate(col_widths):
                    row.cells[idx].width = Inches(w)

        # Space after table
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(6)

    def add_result_text(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.bold = True
            r_pre.font.size = Pt(10.5)
            r_pre.font.color.rgb = COLOR_DARK
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_DARK

    def add_result_image(img_path, width_inches=6.2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run()
        run.add_picture(str(img_path), width=Inches(width_inches))

    # --- Title ---
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("Assignment Tasks & Solutions")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    # ==========================================
    # Task 1
    # ==========================================
    add_heading_1("1. Load the data")
    add_instruction_block("Load the CSV files with Polars and measure the loading time.")

    # Result: Polars loading time summary table
    task1_headers = ["File", "Rows", "Columns", "Size (MB)", "Load Time (s)"]
    task1_rows = [
        ["calendar.csv", "1,969", "13", "0.11", "0.0040"],
        ["sales_test_evaluation.csv", "30,490", "33", "3.11", "0.0212"],
        ["sales_test_validation.csv", "30,490", "33", "2.82", "0.0183"],
        ["sales_train_evaluation.csv", "30,490", "1,946", "115.56", "0.7772"],
        ["sales_train_validation.csv", "30,490", "1,918", "113.91", "0.7164"],
        ["sample_submission.csv", "60,980", "29", "4.99", "0.0422"],
        ["sell_prices.csv", "6,841,121", "4", "226.09", "0.1942"],
        ["weights_evaluation.csv", "42,840", "5", "2.15", "0.0119"],
        ["weights_validation.csv", "42,840", "5", "2.14", "0.0124"]
    ]
    add_result_table(task1_headers, task1_rows, right_align_indices=[1, 2, 3, 4], col_widths=[2.3, 0.9, 0.9, 1.0, 1.1])

    # ==========================================
    # Task 2
    # ==========================================
    add_heading_1("2. Inspect and combine the sales files")
    add_instruction_block(
        "Inspect the sales_train_* and sales_test_* dataframes. There are four of them, described in the ReadMe. "
        "They contain index columns and one column per day, labelled d_{i}. Day d_1 is 29 January 2011.\n\n"
        "Join the data into a single dataframe.\n\n"
        "Replace the integer day columns (d_1, d_2, \u2026) with real dates.\n\n"
        "Persist the result as one Parquet file.\n\n"
        "Tip: check the ReadMe for which days each file covers, so that no day is counted twice."
    )

    # Result: Inspection Table
    task2_inspect_headers = ["File", "Rows", "Total Cols", "Index Columns", "Day Cols", "Day Range"]
    task2_inspect_rows = [
        ["sales_train_validation.csv", "30,490", "1,918", "item_id, dept_id, cat_id, store_id, state_id", "1,913", "d_1 to d_1913"],
        ["sales_test_validation.csv", "30,490", "33", "item_id, dept_id, cat_id, store_id, state_id", "28", "d_1914 to d_1941"],
        ["sales_train_evaluation.csv", "30,490", "1,946", "item_id, dept_id, cat_id, store_id, state_id", "1,941", "d_1 to d_1941"],
        ["sales_test_evaluation.csv", "30,490", "33", "item_id, dept_id, cat_id, store_id, state_id", "28", "d_1942 to d_1969"]
    ]
    add_result_table(task2_inspect_headers, task2_inspect_rows, right_align_indices=[1, 2, 4], col_widths=[1.8, 0.7, 0.7, 1.5, 0.6, 1.0])

    add_result_text("Identical", bold_prefix="Overlap d_1..d_1913 (sales_train_validation vs sales_train_evaluation): ")
    add_result_text("Identical", bold_prefix="Overlap d_1914..d_1941 (sales_test_validation vs sales_train_evaluation): ")
    add_result_text("d_1942 to d_1969 (28 additional days)", bold_prefix="sales_test_evaluation: ")
    add_result_text("30,490 rows \u00d7 1,974 columns (5 index columns + 1,969 day columns from d_1 to d_1969)", bold_prefix="Combined DataFrame shape: ")
    add_result_text("2011-01-29 (Day d_1)", bold_prefix="First date column: ")
    add_result_text("2016-06-19 (Day d_1969)", bold_prefix="Last date column: ")

    # Dated DataFrame sample table
    task2_sample_headers = ["item_id", "dept_id", "cat_id", "store_id", "state_id", "2011-01-29", "2011-01-30", "...", "2016-06-18", "2016-06-19"]
    task2_sample_rows = [
        ["HOBBIES_1_001", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "0", "0", "...", "3", "1"],
        ["HOBBIES_1_002", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "0", "0", "...", "1", "0"],
        ["HOBBIES_1_003", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "0", "0", "...", "1", "1"]
    ]
    add_result_table(task2_sample_headers, task2_sample_rows, right_align_indices=[5, 6, 7, 8, 9], col_widths=[1.1, 0.8, 0.7, 0.6, 0.5, 0.7, 0.7, 0.4, 0.7, 0.7])

    add_result_text("sales_combined.parquet (20.07 MB)", bold_prefix="Persisted Parquet file: ")

    # ==========================================
    # Task 3
    # ==========================================
    add_heading_1("3. From wide to long")
    add_instruction_block(
        "Use Polars unpivot to bring the dataframe from wide format (one column per day) to long format (one row per series and day).\n\n"
        "Tip: if you run into memory errors, use a Polars lazy dataframe (scan_* / .lazy() and .collect())."
    )

    add_result_text("60,034,810 rows (30,490 series \u00d7 1,969 days) \u00d7 7 columns", bold_prefix="Unpivoted DataFrame size: ")

    task3_headers = ["item_id", "dept_id", "cat_id", "store_id", "state_id", "date", "sales"]
    task3_rows = [
        ["HOBBIES_1_001", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0"],
        ["HOBBIES_1_002", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0"],
        ["HOBBIES_1_003", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0"],
        ["HOBBIES_1_004", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0"],
        ["HOBBIES_1_005", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0"]
    ]
    add_result_table(task3_headers, task3_rows, right_align_indices=[5, 6], col_widths=[1.2, 1.0, 0.9, 0.8, 0.7, 1.0, 0.6])

    # ==========================================
    # Task 4
    # ==========================================
    add_heading_1("4. Add prices and revenue")
    add_instruction_block(
        "Join the long dataframe with a subset of the calendar columns and with the sell prices, "
        "so that next to the number of items sold you also have the price. Multiply the two to get the revenue of each transaction."
    )

    task4_headers = ["item_id", "dept_id", "cat_id", "store_id", "state_id", "date", "sales", "wm_yr_wk", "year", "month", "sell_price", "revenue"]
    task4_rows = [
        ["HOBBIES_1_001", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0", "11101", "2011", "1", "null", "0.0"],
        ["HOBBIES_1_004", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0", "11101", "2011", "1", "null", "0.0"],
        ["HOBBIES_1_002", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0", "11101", "2011", "1", "null", "0.0"],
        ["HOBBIES_1_003", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0", "11101", "2011", "1", "null", "0.0"],
        ["HOBBIES_1_005", "HOBBIES_1", "HOBBIES", "CA_1", "CA", "2011-01-29", "0", "11101", "2011", "1", "null", "0.0"]
    ]
    add_result_table(task4_headers, task4_rows, right_align_indices=[6, 7, 8, 9, 10, 11], col_widths=[1.0, 0.8, 0.6, 0.5, 0.4, 0.7, 0.4, 0.5, 0.4, 0.4, 0.5, 0.5])

    # ==========================================
    # Task 5
    # ==========================================
    add_heading_1("5. Calculate")
    add_instruction_block(
        "The total revenue over the whole period (about 5 years).\n\n"
        "Aggregate the revenue by day, month and year. For each of the three aggregations report:\n\n"
        "the average revenue,\n\n"
        "the date and value of the lowest revenue,\n\n"
        "the date and value of the highest revenue."
    )

    add_result_text("$195,374,066.67", bold_prefix="Total Revenue over the whole period: ")

    task5_headers = ["Aggregation", "Average Revenue", "Lowest Revenue", "Highest Revenue"]
    task5_rows = [
        ["Daily", "$99,225.02", "$23.68 (2011-12-25)", "$191,941.73 (2016-05-15)"],
        ["Monthly", "$2,960,213.13", "$218,328.09 (2011-01)", "$4,296,651.61 (2016-05)"],
        ["Yearly", "$32,562,344.44", "$22,631,786.64 (2016)", "$42,416,456.61 (2015)"]
    ]
    add_result_table(task5_headers, task5_rows, right_align_indices=[1, 2, 3], col_widths=[1.2, 1.4, 1.8, 1.8])

    # ==========================================
    # Task 6
    # ==========================================
    add_heading_1("6. Visualise")
    add_instruction_block(
        "Make three plots of the total revenue:\n\n"
        "daily\n\n"
        "monthly\n\n"
        "yearly"
    )

    img_dir = Path("tis/hw1/extracted_assets")
    add_result_image(img_dir / "daily_plot.png", width_inches=6.2)
    add_result_image(img_dir / "monthly_plot.png", width_inches=6.2)
    add_result_image(img_dir / "yearly_plot.png", width_inches=5.2)

    doc.save(output_docx_path)
    print(f"Document successfully created at: {output_docx_path}")

if __name__ == "__main__":
    create_solution_document("tis/hw1/assignment_tasks_and_solutions.docx")
